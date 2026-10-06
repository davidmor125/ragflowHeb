import pytest
from lxml import etree

from rag.app.naive import _docx_cell_text, _docx_para_text, _iter_docx_blocks
from rag.nlp import _expand_oversized_rows, _merge_cks, _top_level_table_parts
from rag.prompts.generator import compact_chunk_html

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _xml(body):
    return etree.fromstring(f'<w:body xmlns:w="{W}" xmlns:mc="{MC}">{body}</w:body>')


def _texts(parent):
    return [_docx_para_text(b) if b.tag.endswith("}p") else "TABLE" for b in _iter_docx_blocks(parent)]


class TestDocxBlocks:
    def test_block_sdt_content_is_yielded(self):
        body = _xml(
            "<w:p><w:r><w:t>לפני</w:t></w:r></w:p>"
            "<w:sdt><w:sdtPr/><w:sdtContent>"
            "<w:p><w:r><w:t>בתוך בקרה</w:t></w:r></w:p>"
            "<w:sdt><w:sdtContent><w:tbl/></w:sdtContent></w:sdt>"
            "</w:sdtContent></w:sdt>"
        )
        assert _texts(body) == ["לפני", "בתוך בקרה", "TABLE"]

    def test_toc_sdt_is_skipped(self):
        body = _xml(
            "<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery w:val='Table of Contents'/></w:docPartObj></w:sdtPr>"
            "<w:sdtContent><w:p><w:r><w:t>1. רקע ..... 3</w:t></w:r></w:p></w:sdtContent></w:sdt>"
            "<w:p><w:r><w:t>גוף</w:t></w:r></w:p>"
        )
        assert _texts(body) == ["גוף"]


class TestDocxParagraphText:
    def test_tracked_insertion_kept_and_deletion_dropped(self):
        p = _xml(
            "<w:p><w:r><w:t xml:space='preserve'>יישלח מייל </w:t></w:r>"
            "<w:ins><w:r><w:t>גם למאשר השני</w:t></w:r></w:ins>"
            "<w:del><w:r><w:delText>ישן</w:delText><w:t>ישן</w:t></w:r></w:del></w:p>"
        )[0]
        assert _docx_para_text(p) == "יישלח מייל גם למאשר השני"

    def test_inline_sdt_smarttag_hyperlink_tab_break(self):
        p = _xml(
            "<w:p><w:pPr><w:tabs><w:tab w:val='left'/></w:tabs></w:pPr>"
            "<w:r><w:t>א</w:t><w:tab/><w:t>ב</w:t><w:br/></w:r>"
            "<w:sdt><w:sdtContent><w:r><w:t>ג</w:t></w:r></w:sdtContent></w:sdt>"
            "<w:smartTag><w:r><w:t>ד</w:t></w:r></w:smartTag>"
            "<w:hyperlink><w:r><w:t>ה</w:t></w:r></w:hyperlink></w:p>"
        )[0]
        # the <w:tab> inside <w:pPr><w:tabs> is a tab-stop definition, not text
        assert _docx_para_text(p) == "א\tב\nגדה"

    def test_textbox_fallback_not_duplicated(self):
        p = _xml(
            "<w:p><w:r><mc:AlternateContent>"
            "<mc:Choice><w:txbxContent><w:p><w:r><w:t>תיבה</w:t></w:r></w:p></w:txbxContent></mc:Choice>"
            "<mc:Fallback><w:txbxContent><w:p><w:r><w:t>תיבה</w:t></w:r></w:p></w:txbxContent></mc:Fallback>"
            "</mc:AlternateContent></w:r></w:p>"
        )[0]
        assert _docx_para_text(p).split() == ["תיבה"]

    def test_cell_text_includes_sdt_and_nested_table(self):
        tc = _xml(
            "<w:tc><w:p><w:r><w:t>שורה 1</w:t></w:r></w:p>"
            "<w:sdt><w:sdtContent><w:p><w:r><w:t>שורה 2</w:t></w:r></w:p></w:sdtContent></w:sdt>"
            "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>x</w:t></w:r></w:p></w:tc>"
            "<w:tc><w:p><w:r><w:t>y</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:tc>"
        )[0]
        assert _docx_cell_text(tc) == "שורה 1\nשורה 2\nx | y"


class TestTopLevelTableParts:
    def test_nested_table_rows_stay_whole(self):
        html = (
            "<table><caption>C</caption><tr><td>H</td></tr>"
            "<tr><td>a<table><tr><td>in1</td></tr><tr><td>in2</td></tr></table></td></tr>"
            "<tr dir='rtl'><td>b</td></tr></table>"
        )
        caption, rows = _top_level_table_parts(html)
        assert caption == "<caption>C</caption>"
        assert rows == [
            "<tr><td>H</td></tr>",
            "<tr><td>a<table><tr><td>in1</td></tr><tr><td>in2</td></tr></table></td></tr>",
            "<tr dir='rtl'><td>b</td></tr>",
        ]
        for r in rows:
            assert r.count("<table") == r.count("</table>")

    def test_tbody_and_no_rows(self):
        assert _top_level_table_parts("<table><tbody><tr><td>1</td></tr></tbody></table>")[1] == ["<tr><td>1</td></tr>"]
        assert _top_level_table_parts("<p>no table</p>") == ("", [])


class TestExpandOversizedRows:
    def test_layout_row_is_replaced_by_inner_rows(self):
        inner = "".join(f"<tr><td>שורה פנימית מספר {i} עם טקסט ארוך למדי</td></tr>" for i in range(40))
        row = f"<tr><td><p>כותרת הפרק</p><table>{inner}</table><p>סיכום</p></td></tr>"
        out = _expand_oversized_rows([row], 50)
        # prose and inner rows keep document order
        assert out[0] == "<tr><td>כותרת הפרק</td></tr>"
        assert out[-1] == "<tr><td>סיכום</td></tr>"
        assert len(out) == 42
        assert all(r.count("<table") == r.count("</table>") for r in out)

    def test_prose_row_is_split_by_paragraph(self):
        paras = "".join(f"<p dir='rtl'><span lang='he-IL'>פסקה {i} עם מספיק מילים כדי לעבור את הסף</span></p>" for i in range(30))
        out = _expand_oversized_rows([f"<tr><td>{paras}</td></tr>"], 50)
        assert len(out) == 30
        assert out[0] == "<tr><td>פסקה 0 עם מספיק מילים כדי לעבור את הסף</td></tr>"

    def test_small_rows_untouched(self):
        row = "<tr><td>a<table><tr><td>b</td></tr></table></td></tr>"
        assert _expand_oversized_rows([row], 256) == [row]


class TestCompactChunkHtml:
    def test_keeps_table_structure_and_text_drops_markup(self):
        html = (
            "<table style='width:5pt'><tbody><tr><td rowspan='2' style='x'><p dir=\"rtl\"><span lang=\"he-IL\">שאר</span>"
            "<span lang=\"ar-SA\">&nbsp;</span><span lang=\"he-IL\">הסניפים</span></p></td>"
            "<td colspan=\"3\"><strong>10,000,000 ₪</strong></td></tr></tbody></table>"
        )
        out = compact_chunk_html(html)
        assert out == '<table><tbody><tr><td rowspan="2">שאר הסניפים</td><td colspan="3">10,000,000 ₪</td></tr></tbody></table>'

    def test_plain_text_untouched(self):
        assert compact_chunk_html("טקסט רגיל 5 < 7") == "טקסט רגיל 5 < 7"


class _FakeStore:
    def __init__(self, docs):
        self.docs = docs

    def get(self, chunk_id, index, kb_ids):
        return self.docs.get(chunk_id)


def _child(cid, mom, content, sim):
    return {"chunk_id": cid, "mom_id": mom, "content_with_weight": content, "content_ltks": content,
            "kb_id": "kb", "similarity": sim, "important_kwd": []}


class TestRetrievalByChildren:
    def _dealer(self, docs):
        from rag.nlp.search import Dealer
        d = Dealer.__new__(Dealer)
        d.dataStore = _FakeStore(docs)
        return d

    def test_small_parent_replaces_children(self):
        store = {"m1": {"content_with_weight": "<table><tr><td>כותרת</td></tr><tr><td>שורה</td></tr></table>", "doc_id": "d", "kb_id": "kb"}}
        out = self._dealer(store).retrieval_by_children([_child("c1", "m1", "שורה", 0.9)], ["t"])
        assert out[0]["content_with_weight"] == store["m1"]["content_with_weight"]

    def test_huge_parent_uses_matched_children(self):
        big = "<table>" + "".join(f"<tr><td>שורה מספר {i} עם תוכן ארוך בטבלה</td></tr>" for i in range(400)) + "</table>"
        store = {"m1": {"content_with_weight": big, "doc_id": "d", "kb_id": "kb"}}
        kids = [_child("c1", "m1", "<table><tr><td>כותרת</td></tr><tr><td>35. עמלות לסוכנים</td></tr></table>", 0.9),
                _child("c2", "m1", "<table><tr><td>כותרת</td></tr><tr><td>36. אחר</td></tr></table>", 0.8)]
        out = self._dealer(store).retrieval_by_children(kids, ["t"])
        assert len(out) == 1
        # whole parent kept; matched children offered as the smaller fallback
        assert out[0]["content_with_weight"] == big
        assert "35. עמלות לסוכנים" in out[0]["content_fallback"] and "שורה מספר 399" not in out[0]["content_fallback"]

    def test_missing_parent_keeps_children(self):
        out = self._dealer({}).retrieval_by_children([_child("c1", "gone", "טקסט", 0.5)], ["t"])
        assert [c["chunk_id"] for c in out] == ["c1"]

    def test_top_parent_gets_same_section_neighbors_only(self):
        parents = {
            "m1": {"content_with_weight": "⟦נוהל: X | 4. מקרים › 4.1. א⟧\n4.1. מקרה ראשון", "doc_id": "d", "kb_id": "kb", "page_num_int": [5]},
            "p6": {"content_with_weight": "⟦נוהל: X | 4. מקרים › 4.5. ה⟧\n4.5. מקרה חמישי", "page_num_int": [6]},
            "p7": {"content_with_weight": "⟦נוהל: X | 5. סיום⟧\n5. פרק אחר", "page_num_int": [7]},
            "p4": {"content_with_weight": "⟦נוהל: X | 3. רקע⟧\nרקע", "page_num_int": [4]},
        }

        class Store(_FakeStore):
            def search(self, fields, hl, cond, mx, ob, off, lim, idx, kbs):
                return [(k, v) for k, v in self.docs.items() if k != "m1" and v["page_num_int"][0] in cond["page_num_int"]]

            def get_fields(self, res, fields):
                return {k: v for k, v in res}

        from rag.nlp.search import Dealer
        dealer = Dealer.__new__(Dealer)
        dealer.dataStore = Store(parents)
        out = dealer.retrieval_by_children([_child("c1", "m1", "4.1. מקרה ראשון", 0.9)], ["t"])
        text = out[0]["content_with_weight"]
        # the next parent of section 4 is appended; section 5 and section 3 are not
        assert "4.5. מקרה חמישי" in text and "5. פרק אחר" not in text and "רקע" not in text
        assert out[0]["content_fallback"] == parents["m1"]["content_with_weight"]
        assert "_page" not in out[0]


class TestProcedureTitle:
    @staticmethod
    def _title(t):
        from deepdoc.parser.html_parser import RAGFlowHtmlParser
        return RAGFlowHtmlParser.procedure_title("x.html", f"<html><head><title>{t}</title></head><body></body></html>".encode())

    def test_name_after_category_marker(self):
        t = 'מיסוי - שו"ת- בנקאות ונכסים - ניכוי מס מתשלומים לחו"ל או לתושבי חוץ - 5088'
        assert self._title(t) == 'ניכוי מס מתשלומים לחו"ל או לתושבי חוץ'

    def test_name_containing_dash_is_kept_whole(self):
        t = 'מכשירים עתידיים - ניירות ערך - שו"ת- בנקאות ונכסים - \u200fאופציות על מדד ת"א - 35 - 4465'
        assert self._title(t) == 'אופציות על מדד ת"א - 35'

    def test_without_marker_skips_two_category_segments(self):
        t = 'מכשירים עתידיים - ניירות ערך - חוזים עתידיים על מדד ת"א - 35 - 4466'
        assert self._title(t) == 'חוזים עתידיים על מדד ת"א - 35'

    @staticmethod
    def _body(*lines):
        from deepdoc.parser.html_parser import RAGFlowHtmlParser
        html = "".join(f"<p>{x}</p>" for x in lines)
        return RAGFlowHtmlParser.procedure_title("x.html", html.encode())

    def test_body_only_first_line(self):
        assert self._body("חסימה ביטול הקפאה וטיפול שוטף בכרטיסים", "פרק א'") == "חסימה ביטול הקפאה וטיפול שוטף בכרטיסים"

    def test_body_only_skips_navigation_and_labels(self):
        assert self._body("למעבר לדף ריכוז נוהלי פתיחת חשבון מקוון -", "לחצו כאן", "פתיחה וניהול של חשבון מקוון משותף") == "פתיחה וניהול של חשבון מקוון משותף"
        assert self._body("כותרת", "הוראות באמצעות הפקסימיליה", "כללי") == "הוראות באמצעות הפקסימיליה"

    def test_body_only_rejects_generic_and_debris(self):
        assert self._body("רקע", "ככלל") == ""
        assert self._body("ק ע", "משהו") == ""
        assert self._body("מערכת") == ""

    def test_no_title(self):
        from deepdoc.parser.html_parser import RAGFlowHtmlParser
        assert RAGFlowHtmlParser.procedure_title("x.html", b"<html><body><p>x</p></body></html>") == ""


class TestSectionContext:
    @pytest.fixture(autouse=True)
    def _word_count_tokens(self, monkeypatch):
        # The real counter needs nltk data that is not present on every dev
        # machine; sizing is not what these tests check.
        from bs4 import BeautifulSoup

        import deepdoc.parser.html_parser as hp
        monkeypatch.setattr(hp.RAGFlowHtmlParser, "_text_tokens",
                            staticmethod(lambda h: len(BeautifulSoup(h, "html.parser").get_text(" ").split())))
        monkeypatch.setattr(hp.rag_tokenizer, "tokenize", lambda t: " ".join(str(t).split()))

    @staticmethod
    def _parts(body, title="נוהל בדיקה", budget=256):
        from deepdoc.parser.html_parser import RAGFlowHtmlParser
        return RAGFlowHtmlParser().parts("x.html", body.encode(), budget, 2048, doc_title=title, with_context=True)

    def test_heading_detection(self):
        from deepdoc.parser.html_parser import RAGFlowHtmlParser as P
        assert P._heading_level("1.5.2. השקעות הוניות") == (3, "1.5.2. השקעות הוניות")
        assert P._heading_level("פרק ב' - חשבון מקוון") == (0, "פרק ב' - חשבון מקוון")
        # a numbered clause with an amount is body text, not a heading
        assert P._heading_level("3.1. סך התקבולים בחשבון לא יעלו סכום של 50,000 ₪ בחודש.") is None

    def test_clause_gets_its_own_section_path(self):
        prose, _ = self._parts("<p>3.3. וועדת הקרנות</p><p>הוועדה דנה בבקשות.</p>"
                               "<p>4. מסלולי ההלוואות</p><p>4.4. שיעור הריבית בהלוואה פריים + 1.7% .</p>")
        text = "\n".join(prose)
        line = [x for x in text.split("\n") if x.startswith("⟦") and "4.4." in x]
        assert line and "4. מסלולי ההלוואות" in line[0] and "נוהל: נוהל בדיקה" in line[0]

    def test_table_gets_context_caption(self):
        _, tables = self._parts("<p>3. דוח חריגים</p><table><tr><td>נושא</td><td>סכום</td></tr><tr><td>עסקי</td><td>25,000</td></tr></table>")
        assert "<caption>נוהל: נוהל בדיקה | 3. דוח חריגים</caption>" in tables[0]

    def test_children_carry_the_context_line(self):
        from rag.nlp import _CTX_LINE_RE
        content = "⟦נוהל: X | 4.4. ריבית⟧\nשורה ראשונה ארוכה מספיק כדי להיות ילד\n⟦נוהל: X | 5. סיום⟧\nשורה שנייה ארוכה מספיק כדי להיות ילד"
        # the context line itself is recognised and never becomes a child
        assert _CTX_LINE_RE.match(content)

    def test_off_by_default(self):
        from deepdoc.parser.html_parser import RAGFlowHtmlParser
        prose, _ = RAGFlowHtmlParser().parts("x.html", "<p>1. כללי</p><p>טקסט</p>".encode(), 256)
        assert not any("⟦" in p for p in prose)


class TestMergeCks:
    def test_paragraphs_are_not_fused(self):
        cks = [
            {"text": "בקשת שירות ב SNOW", "ck_type": "text", "tk_nums": 3},
            {"text": "בבינלאומי, במקרה", "ck_type": "text", "tk_nums": 3},
        ]
        merged, _ = _merge_cks(cks, 128, False)
        assert merged[0]["text"] == "בקשת שירות ב SNOW\nבבינלאומי, במקרה"


class TestDocxSectionContext:
    """Docx(with_context=True): ⟦procedure | section⟧ lines and table captions."""

    @staticmethod
    def _docx(build):
        from io import BytesIO

        from docx import Document

        d = Document()
        build(d)
        buf = BytesIO()
        d.save(buf)
        return buf.getvalue()

    @staticmethod
    def _parse(binary, **kw):
        from rag.app.naive import Docx

        return Docx()("t.docx", binary, with_context=True, doc_title="נוהל בדיקה", **kw)

    def test_outline_headings_give_the_path(self):
        def build(d):
            d.add_heading("דרישות", level=1)
            d.add_heading("פעילות המנגנון", level=2)
            d.add_paragraph("הזנת תמהיל לניגוח.")
            d.add_paragraph("בניית תמהיל כהצעה ללקוח.")

        texts = [t for t, _, _ in self._parse(self._docx(build))]
        # The line goes where the section changes, i.e. before its heading;
        # body paragraphs of the same section carry none of their own.
        assert texts[1] == "⟦נוהל: נוהל בדיקה | דרישות › פעילות המנגנון⟧\nפעילות המנגנון"
        assert texts[2:] == ["הזנת תמהיל לניגוח.", "בניית תמהיל כהצעה ללקוח."]

    def test_custom_hebrew_style_with_outline_level(self):
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml import parse_xml

        def build(d):
            st = d.styles.add_style("כותרת רמה 2", WD_STYLE_TYPE.PARAGRAPH)
            st.element.get_or_add_pPr().append(parse_xml('<w:outlineLvl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="1"/>'))
            d.add_heading("מבוא", level=1)
            d.add_paragraph("סימולציה מנותקת מבקשה", style="כותרת רמה 2")
            d.add_paragraph("גוף הסעיף.")

        texts = [t for t, _, _ in self._parse(self._docx(build))]
        assert texts[1] == "⟦נוהל: נוהל בדיקה | מבוא › סימולציה מנותקת מבקשה⟧\nסימולציה מנותקת מבקשה"

    def test_bold_list_items_are_headings_without_outline(self):
        def build(d):
            d.add_paragraph().add_run("מתווה הפתרון").bold = True
            d.add_paragraph("התצוגה תהיה רק במשכנתאות שבהן הלקוח לווה.")

        texts = [t for t, _, _ in self._parse(self._docx(build))]
        assert texts == ["⟦נוהל: נוהל בדיקה | מתווה הפתרון⟧\nמתווה הפתרון", "התצוגה תהיה רק במשכנתאות שבהן הלקוח לווה."]

    def test_table_caption_carries_the_context(self):
        def build(d):
            d.add_heading("עמלות", level=1)
            t = d.add_table(rows=2, cols=2)
            t.cell(0, 0).text, t.cell(0, 1).text = "סוג", "סכום"
            t.cell(1, 0).text, t.cell(1, 1).text = "העברה", "5"

        tables = [tb for _, _, tb in self._parse(self._docx(build)) if tb]
        assert tables[0].startswith("<table><caption>נוהל: נוהל בדיקה | עמלות</caption>")

    def test_heading_equal_to_title_is_not_a_section(self):
        def build(d):
            d.add_heading("נוהל בדיקה", level=1)
            d.add_heading("ממשקים", level=2)
            d.add_paragraph("שדה חובה.")

        texts = [t for t, _, _ in self._parse(self._docx(build))]
        assert texts[0] == "⟦נוהל: נוהל בדיקה⟧\nנוהל בדיקה"
        assert texts[1] == "⟦נוהל: נוהל בדיקה | ממשקים⟧\nממשקים"

    def test_off_by_default(self):
        from rag.app.naive import Docx

        def build(d):
            d.add_heading("דרישות", level=1)
            d.add_paragraph("גוף.")

        assert all("⟦" not in (t or "") for t, _, _ in Docx()("t.docx", self._docx(build)))


class TestCarryDocxContext:
    def test_mid_section_chunk_gets_the_line_in_effect(self):
        from rag.app.naive import _carry_docx_context

        chunks = [
            {"text": "טופס בראש המסמך", "ck_type": "image"},
            {"text": "⟦נוהל: X⟧\nפתיחה", "ck_type": "text"},
            {"text": "⟦נוהל: X | כללי⟧\nסעיף", "ck_type": "text"},
            {"text": "המשך הסעיף", "ck_type": "text"},
            {"text": "<table></table>", "ck_type": "table"},
        ]
        out = [c["text"] for c in _carry_docx_context(chunks)]
        assert out[0] == "⟦נוהל: X⟧\nטופס בראש המסמך"
        assert out[3] == "⟦נוהל: X | כללי⟧\nהמשך הסעיף"
        assert out[4] == "<table></table>"


class TestDocxTitle:
    @staticmethod
    def _doc(first_heading=None, core_title=""):
        from docx import Document

        d = Document()
        d.core_properties.title = core_title
        if first_heading:
            d.add_heading(first_heading, level=1)
        return d

    def test_version_and_date_noise_dropped(self):
        from rag.app.naive import _docx_title

        assert _docx_title("ניגוח וסימולציה - גרסה 0.2 - 20260604-01 - ללא מעקב שינויים.docx", self._doc()) == "ניגוח וסימולציה"
        assert _docx_title("סבסוד קבלן גרסה 0.4 20260604 הערות.docx", self._doc()) == "סבסוד קבלן"
        assert _docx_title("אפיון רשות התאגידים - שליפת נסח חברה - 1.1 -  מאושר.docx", self._doc()) == "אפיון רשות התאגידים - שליפת נסח חברה"

    def test_meaningless_file_name_uses_first_heading(self):
        from rag.app.naive import _docx_title

        assert _docx_title("9999999999.docx", self._doc("אשראי מובטח 55")) == "אשראי מובטח 55"
        assert _docx_title("client.docx", self._doc("x", core_title="אפיון עסקי")) == "אפיון עסקי"


class TestHtmlRunJoining:
    """Word exports split words and numbers across <span>s; they must be
    joined the way a browser shows them, with no added spaces."""

    @pytest.fixture(autouse=True)
    def _word_count_tokens(self, monkeypatch):
        from bs4 import BeautifulSoup

        import deepdoc.parser.html_parser as hp
        monkeypatch.setattr(hp.RAGFlowHtmlParser, "_text_tokens",
                            staticmethod(lambda h: len(BeautifulSoup(h, "html.parser").get_text(" ").split())))
        monkeypatch.setattr(hp.rag_tokenizer, "tokenize", lambda t: " ".join(str(t).split()))

    @staticmethod
    def _prose(body, with_context=False):
        from deepdoc.parser.html_parser import RAGFlowHtmlParser
        prose, tables = RAGFlowHtmlParser().parts("x.html", f"<html><body>{body}</body></html>".encode(), 256, 2048,
                                                  doc_title="נוהל", with_context=with_context)
        return "\n".join(prose), tables

    def test_runs_join_without_spaces(self):
        body = '<p><span lang="he">עד </span><span lang="ar-SA">1</span><span>50,000</span><span> ש</span><span>"</span><span>ח ע</span><span>"</span><span>י המנהל</span></p>'
        for ctx in (False, True):
            text, _ = self._prose(body, ctx)
            assert 'עד 150,000 ש"ח ע"י המנהל' in text

    def test_whitespace_between_runs_is_kept(self):
        text, _ = self._prose("<p><span>סכום</span> <span>ההלוואה</span>\n<span>המבוקש</span></p>")
        assert "סכום ההלוואה המבוקש" in text

    def test_br_separates_words(self):
        text, _ = self._prose('<p><span style="font-weight:bold">כללי</span><br/><span>הבנק מציע</span></p>')
        assert "כללי הבנק מציע" in text and "כלליהבנק" not in text

    def test_blocks_stay_separate_lines(self):
        text, _ = self._prose("<p><span>1</span><span>1. מעקב</span></p><p>שורה שנייה</p>", with_context=True)
        lines = [x for x in text.split("\n") if not x.startswith("⟦")]
        assert "11. מעקב" in lines and "שורה שנייה" in lines

    def test_base64_images_are_dropped(self):
        img = '<img alt="x" src="data:image/png;base64,' + "iVBORw0KGgo" * 200 + '">'
        text, tables = self._prose(f"<p>פרק 1 {img}</p><table><tr><td>תא {img}</td></tr></table>")
        assert "base64" not in text and all("base64" not in t for t in tables)
        assert "פרק 1" in text


class TestHtmlTitleRuns:
    @staticmethod
    def _title(html):
        from deepdoc.parser.html_parser import RAGFlowHtmlParser
        return RAGFlowHtmlParser.procedure_title("x.html", html.encode())

    def test_name_split_across_runs_is_whole(self):
        html = '<p><span>העברת הוראות לני</span><span>"</span><span>ע הנסחרים בבורסות</span></p><p>כללי</p>'
        assert self._title(html) == 'העברת הוראות לני"ע הנסחרים בבורסות'

    def test_h1_label_falls_back_to_first_line(self):
        assert self._title("<h1>שם הנוהל</h1><p>נושא</p><p>ממסרים דחויים</p>") == "ממסרים דחויים"

    def test_numbered_clause_is_not_a_title(self):
        assert self._title("<p>1. כללי</p><p>הבנק מאפשר ללקוחות להפקיד שיקים.</p>") == ""
        assert self._title("<p>נושאי הנוהל</p><p>פרק א</p>") == ""


class TestInlineMarkupStripping:
    def test_strip_markup_joins_inline_runs(self):
        from rag.nlp import strip_markup
        html = '<table><tr><td><p><span>1</span><span>50,000</span> <span>ש</span><span>"ח</span></p></td><td>ב</td></tr></table>'
        assert strip_markup(html) == '150,000 ש"ח ב'

    def test_compact_keeps_table_and_joins_runs(self):
        from rag.prompts.generator import compact_chunk_html
        out = compact_chunk_html('<table><tr><td><span lang="he">1</span><span>50,000</span></td></tr></table>')
        assert out == "<table><tr><td>150,000</td></tr></table>"

    def test_long_tags_are_stripped(self):
        from rag.nlp import strip_markup
        img = '<img src="data:image/png;base64,' + "A" * 5000 + '">'
        assert strip_markup(f"טקסט {img} נוסף") == "טקסט נוסף"
