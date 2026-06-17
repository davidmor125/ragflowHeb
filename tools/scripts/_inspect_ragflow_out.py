"""Look at the RAGFlow eval output."""
import sys
import openpyxl
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

p = Path(r"C:/develop/html_output/_eval_report/newtest/test_report_24_05_ragflow.xlsx")
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
ws = wb.active

# 1-based: nohel=B=2, question=E=5, answer=F=6, gold=G=7, link=H=8, verdict=J=10
for r in range(6, 17):
    nohel = ws.cell(row=r, column=2).value
    answer = ws.cell(row=r, column=6).value
    gold = ws.cell(row=r, column=7).value
    link = ws.cell(row=r, column=8).value
    verdict = ws.cell(row=r, column=10).value
    if not nohel:
        continue
    print(f"=== ROW {r} (נוהל={nohel}) ===")
    q = ws.cell(row=r, column=5).value
    print(f"  שאלה:        {(str(q)[:90] if q else '')}")
    print(f"  זהב:         {(str(gold)[:90] if gold else '')}")
    print(f"  תשובת מערכת: {(str(answer)[:200] if answer else '<EMPTY>')}")
    print(f"  link:        {link!r}")
    print(f"  verdict:     {(str(verdict)[:120] if verdict else '<EMPTY>')}")
    print()
